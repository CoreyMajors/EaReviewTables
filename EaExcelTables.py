import openpyxl
from openpyxl import load_workbook
import xlsxwriter


# Create an new Excel file and add a worksheet.
workbook = xlsxwriter.Workbook('EARtables.xlsx')

#overview
worksheet = workbook.add_worksheet("Overview")

worksheet.add_table('A1:N100', {'banded_columns': True})

worksheet.write('A1','Initative Name')
worksheet.write('B1', 'Description')
worksheet.write('C1', 'Start Date')
worksheet.write('D1', 'End Date')
worksheet.write('E1', 'Project Sponsor')
worksheet.write('F1', 'Product Owner')
worksheet.write('G1', 'Capability Leader')
worksheet.write('H1', 'Solution Leader')
worksheet.write('I1', 'Application Owner*')
worksheet.write('J1', 'Requesting Org')
worksheet.write('K1', 'Investment Category')
worksheet.write('L1', 'Total Investment')
worksheet.write('M1', 'T-Shirt Size')
worksheet.write('N1', 'MT Project Manager | Scrum Master')


#BCM
worksheet1 = workbook.add_worksheet("Business Capabilties")

worksheet1.add_table("A1:S100", {'banded_columns': True})

worksheet1.write('A1','Initative Name')
worksheet1.write('B1', '1) Which processes in the BCM does the project support? Level 1')
worksheet1.write('C1', '1) Which processes in the BCM does the project support? Level 2')
worksheet1.write('D1', '1) Which processes in the BCM does the project support? Level 3')
worksheet1.write('E1', '1) Which processes in the BCM does the project support? Level 4')
worksheet1.write('F1', '2) Which processes in the BCM does the project support? Level 1')
worksheet1.write('G1', '2) Which processes in the BCM does the project support? Level 2')
worksheet1.write('H1', '2) Which processes in the BCM does the project support? Level 3')
worksheet1.write('I1', '2) Which processes in the BCM does the project support? Level 4')
worksheet1.write('J1', '3) Which processes in the BCM does the project support? Level 1')
worksheet1.write('K1', '3) Which processes in the BCM does the project support? Level 2')
worksheet1.write('L1', '3) Which processes in the BCM does the project support? Level 3')
worksheet1.write('M1', '3) Which processes in the BCM does the project support? Level 4')
worksheet1.write('N1', 'Does the solution or initiative support more than 3 capabilities?')
worksheet1.write('O1', 'If yes, provide a list of the others')
worksheet1.write('P1', 'Does this enable new business capabilities for McKesson?')
worksheet1.write('Q1', 'If yes, provide a description')
worksheet1.write('R1', 'Does this align with the capability leader�s technology road map?')
worksheet1.write('S1', 'If no, please explain')


#applicaions
worksheet2 = workbook.add_worksheet("Applications")

worksheet2.add_table("A1:E100", {'banded_columns': True})

worksheet2.write('A1','Initative Name')
worksheet2.write('B1', 'Have any existing solutions in our landscape been explored?')
worksheet2.write('C1', 'Will this solution be external-facing?')
worksheet2.write('D1', 'What is the LeanIX ID for the application?')
worksheet2.write('E1', 'Has a conceptual architecture diagram been created?')



#Integration
worksheet3 = workbook.add_worksheet("Integration")

worksheet3.add_table("A1:E100", {'banded_columns': True})

worksheet3.write('A1','Initative Name')
worksheet3.write('B1', 'What systems will this solution integrate with?')
worksheet3.write('C1', 'Has a data flow diagram been created?')
worksheet3.write('D1', 'What patterns and platforms will be used for those integrations?')
worksheet3.write('E1', 'What is the volume, frequency and service level requirements for those integrations?')

#security
worksheet4 = workbook.add_worksheet("Security")

worksheet4.add_table("A1:F100", {'banded_columns': True})

worksheet4.write('A1','Initative Name')
worksheet4.write('B1', 'Are there single sign-on requirements?')
worksheet4.write('C1', 'If not, why not?')
worksheet4.write('D1', 'Who from ISRM has been engaged?')
worksheet4.write('E1', 'If new, have the appropriate reviews been completed?')
worksheet4.write('F1', 'Has a risk assessment been completed? ')


#procurement
worksheet5 = workbook.add_worksheet("Procurement")

worksheet5.add_table("A1:C100", {'banded_columns': True})

worksheet5.write('A1','Initative Name')
worksheet5.write('B1', 'What are the licensing requirements?')
worksheet5.write('C1', 'Who from Procurement has been engaged?')

#data
worksheet6 = workbook.add_worksheet("Data")

worksheet6.add_table("A1:D100", {'banded_columns': True})

worksheet6.write('A1','Initative Name')
worksheet6.write('B1', 'What master data elements will be used, modified or created?')
worksheet6.write('C1', 'What are the analytics requirements?')
worksheet6.write('D1', 'What is the approach (policy, process and technology) for information life cycle management?')

#infrastructure
worksheet7 = workbook.add_worksheet("Infrastructure")

worksheet7.add_table("A1:H100", {'banded_columns': True})

worksheet7.write('A1','Initative Name')
worksheet7.write('B1', 'Does the application�s technology stack align with our strategy and standards?')
worksheet7.write('C1', 'Where will the application be hosted?')
worksheet7.write('D1', 'What is the expected volume of network traffic?')
worksheet7.write('E1', 'What is the user count and type?')
worksheet7.write('F1', 'What is the expected growth?')
worksheet7.write('G1', 'What is the cost of downtime?  ')
worksheet7.write('H1', 'Who has been engaged to discuss the appropriate HA and DR capabilities?')

#operations
worksheet8 = workbook.add_worksheet("Operations")

worksheet8.add_table("A1:F100", {'banded_columns': True})

worksheet8.write('A1','Initative Name')
worksheet8.write('B1', 'What groups will be expected to support this solution?')
worksheet8.write('C1', 'Who from those groups has been engaged?')
worksheet8.write('D1', 'Are there IT components that are at risk of EOL')
worksheet8.write('E1', 'If yes, what are they?')
worksheet8.write('F1', 'Does the solution affect multiple regions/BUs?')

#compliance
worksheet9 = workbook.add_worksheet("Compliance")

worksheet9.add_table("A1:F100", {'banded_columns': True})

worksheet9.write('A1','Initative Name')
worksheet9.write('B1', 'What are the compliance requirements?')
worksheet9.write('C1', 'Are there data residency requirements?')
worksheet9.write('D1', 'What regulations apply (e.g., SOx, HIPAA, GDPR)?')
worksheet9.write('E1', 'Is it regulated by an outside entity?')
worksheet9.write('F1', 'Is it legally mandated or contractually required?')

